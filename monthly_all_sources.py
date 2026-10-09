# -*- coding: utf-8 -*-
"""Build the monthly industry workbook from every daily ALL worksheet.

The first week can additionally contribute a translated weekly ALL worksheet.
Those sheets overlap heavily, so only one record per DOI/link/title is kept.
"""
from __future__ import annotations

import re
from calendar import monthrange
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment
from excel_output_utils import format_literature_worksheet
from monthly_industry_view import ROOTS, TAXONOMY, classify_five, _copy_style


HEADERS = ("出版商", "期刊名", "标题", "通讯作者", "发表日期", "DOI", "五级分类", "数量")
TOPICS = ("光伏", "正极", "负极", "电解质", "非活性材料等", "其他储能器件",
          "热能", "核能", "氢氨醇", "氢基能源", "产业降碳", "CCUS", "塑料回收")


def is_climate_policy(title: str) -> bool:
    """Only explicit policy, coalition and political-economy titles match.

    'Regulation' alone is deliberately excluded: it commonly describes
    electrochemical, biological or control-system mechanisms.
    """
    t = title.casefold()
    if re.search(r"\b(policy learning|policy trading operation|political characteristics)\b", t):
        return False
    subject = re.search(
        r"\b(climate|carbon|emissions?|decarboniz\w*|clean technolog\w*|energy|electricity|hydrogen|ccs|industry|industrial|plastics?|green transition)\b", t)
    governance = re.search(
        r"\b(policy|policies|political|coalition|alliance|treaty|carbon pricing|just transition|energy poverty|geopolitical risk|climate agreement)\b", t)
    return bool(subject and governance) or bool(re.search(
        r"\b(building energy performance regulation|carbon border penalties|domestic carbon pricing|climate.disaster funding)\b", t))


POLICY_PATH = ("零碳产业", "物质循环", "资源回收利用与排放治理", "其他环境治理")
GRID_PATH = ("零碳产业", "能量转化", "能量分配与运输", "电网相关技术")


def _industrial_decarbonization(title: str) -> tuple[str, ...]:
    """Disambiguate the broad upstream tag using the paper's actual subject."""
    t = title.casefold()
    if re.search(r"\b(decarbonization pathways|strategic low.carbon transition|financing cost differences|system.level decarbonization|material inequality|urban material inequality)\b", t):
        return POLICY_PATH
    if re.search(r"\b(data cent(?:er|re)s?)\b", t):
        return ("AI与智能科技", "AI硬件层", "数据中心")
    if re.search(r"\b(fishing vessels?|maritime|shipping|ship\b)", t):
        return ("通用技术", "通信和运输", "物质运输", "水路运输")
    if re.search(r"\b(aviation|aircraft)\b", t):
        return ("通用技术", "通信和运输", "物质运输", "航空")
    if re.search(r"\b(vehicles?|robotaxis?|transportation|vehicle routing|freight)\b", t) and not re.search(r"\b(charging|v2g|grid|power.transportation networks?)\b", t):
        return ("通用技术", "通信和运输", "物质运输", "陆路运输")
    if re.search(r"\b(microgrids?|electricity|power systems?|energy systems?|charging|v2g|dispatch|flexible loads?|energy communit\w*|multi.energy flow|grid.interaction)\b", t):
        return GRID_PATH
    if re.search(r"\b(battery storage|battery system technologies|hybrid storage)\b", t):
        return ("零碳产业", "能量转化", "能量存储", "电化学储能", "二次电池")
    if re.search(r"\b(cooling|hvac|thermal insulation|radiative cooling|thermal rectification|thermal.radiative regulation)\b", t):
        return ("零碳产业", "能量转化", "能源测", "二次能源利用", "制冷散热技术")
    if re.search(r"\b(wastewater|remediation|pollution)\b", t):
        return ("零碳产业", "物质循环", "资源回收利用与排放治理", "污水处理" if "wastewater" in t else "其他环境治理")
    if re.search(r"\b(syngas|methanol|petrochemical|acetylene|alkene|chemical manufacturing|adipic acid|propane dehydrogenation|lignin.derived|tetrahydrofurfuryl|methane to methanol)\b", t):
        return ("零碳产业", "物质循环", "资源加工", "有机物", "平台化工品")
    if re.search(r"\b(steel|iron ore|aluminum|cement|concrete)\b", t):
        return ("零碳产业", "物质循环", "资源加工", "无机物", "非金属" if "cement" in t or "concrete" in t else "金属")
    if re.search(r"\b(lignin adhesive|wood bonding|hydrogel|smart windows?|electrochromic|metamaterial reactors?|monolithic film)\b", t):
        return ("通用技术", "材料工程", "其它先进材料")
    if re.search(r"\b(vertical farming|agricultural)\b", t):
        return ("零碳产业", "物质循环", "资源获取", "种植养殖技术")
    if re.search(r"\b(building retrofit|building control|buildings?|industrial park|thermochemical reaction)\b", t):
        return ("通用技术", "工艺和工程")
    return ("通用技术", "工艺和工程")


def source_name(source: str) -> tuple[str, str]:
    text = (source or "").strip()
    if "Nature Communications" in text:
        return "Springer-Nature", "Nature Communications"
    if text.startswith("Nature"):
        return "Springer-Nature", text
    if text.startswith("AAAS: Science Advances") or text == "Science Advances":
        return "AAAS", "Science Advances"
    if text == "Science":
        return "AAAS", "Science"
    if "Chemical Society Reviews" in text or "Chem. Soc. Rev." in text:
        return "RSC", "Chemical Society Reviews"
    if "Energy & Environmental Science" in text or "Energy Environ. Sci." in text:
        return "RSC", "Energy & Environment Science"
    if "Journal of the American Chemical Society" in text:
        return "ACS", "JACS"
    if "ACS Applied Materials & Interfaces" in text:
        return "ACS", "ACS Applied Materials & Interfaces"
    if "ACS Energy Letters" in text:
        return "ACS", "ACS Energy Letters"
    if "Chemical Reviews" in text:
        return "ACS", "Chemical Reviews"
    if "Angewandte Chemie" in text:
        return "Wiley", "Angewandte"
    if "Advanced Energy Materials" in text:
        return "Wiley", "Advanced Energy Materials"
    if "Advanced Materials" in text:
        return "Wiley", "Advanced Materials"
    if text.startswith("ScienceDirect Publication: "):
        return "Elsevier", text.removeprefix("ScienceDirect Publication: ")
    if text in {"Joule", "Chem", "Matter", "One Earth"}:
        return "Elsevier", text
    if "Proceedings of the National Academy" in text:
        return "Others", "PNAS"
    if "Phys. Rev. Lett." in text:
        return "Others", "Physical Review Letters"
    if "Phys. Rev. C" in text:
        return "Others", "Physical Review C"
    if text.startswith("IEEE "):
        return "Others", text.removeprefix("IEEE ")
    if "Nuclear Fusion" in text:
        return "Others", "Nuclear Fusion"
    if "National Science Review" in text:
        return "Others", "National Science Review"
    if "Nano-Micro Letters" in text:
        return "Others", "Nano-Micro Letters"
    return "Others", text or "来源未注明"


def _record_key(row: dict) -> str:
    doi = str(row.get("doi") or "").strip().casefold()
    doi = re.sub(r"^https?://(?:dx\.)?doi\.org/", "", doi)
    if doi.startswith("10."):
        return "doi:" + doi
    link = str(row.get("link") or "").strip().casefold().split("?", 1)[0]
    if link:
        return "url:" + link
    return "title:" + re.sub(r"\s+", " ", str(row.get("title") or "").strip().casefold())


def _date(row: dict):
    for field in ("pub_date", "published", "published_str"):
        value = row.get(field)
        if value:
            parsed = pd.to_datetime(value, errors="coerce", utc=True)
            if not pd.isna(parsed):
                return parsed.date()
    return None


def _read_all(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "ALL" not in wb:
        raise ValueError(f"Missing ALL worksheet: {path}")
    values = wb["ALL"].values
    headers = next(values)
    for values in values:
        yield dict(zip(headers, values))
    wb.close()


def collect(month: str, daily_dir: Path, weekly_dir: Path) -> tuple[list[dict], dict]:
    daily = sorted(daily_dir.glob(f"news_with_abstract_{month}-??_zai_classified.xlsx"))
    weekly = sorted(weekly_dir.glob(f"weekly_news_with_abstract_{month}-??_translated.xlsx"))
    if not daily:
        raise FileNotFoundError(f"No daily Excel workbooks for {month} in {daily_dir}")
    year, mon = map(int, month.split("-"))
    expected = {f"{month}-{day:02d}" for day in range(1, monthrange(year, mon)[1] + 1)}
    observed = {p.name[19:29] for p in daily}
    missing = sorted(expected - observed)
    if missing:
        raise FileNotFoundError(f"Missing daily classified Excel files: {', '.join(missing)}")
    # ALL appears once per workbook; never read its topical copies as new rows.
    files = daily + weekly
    records = {}
    stats = Counter()
    for path in files:
        stats["files"] += 1
        for row in _read_all(path):
            stats["input_rows"] += 1
            date = _date(row)
            if not date or date.strftime("%Y-%m") != month:
                stats["outside_month"] += 1
                continue
            if not str(row.get("title") or "").strip():
                stats["empty_title"] += 1
                continue
            key = _record_key(row)
            if key in records:
                stats["duplicates"] += 1
                old = records[key]
                # Preserve the better metadata when a weekly translated record
                # complements a daily record, without multiplying the article.
                for field in ("abstract", "categories", "last_author", "doi", "link"):
                    if not old.get(field) and row.get(field):
                        old[field] = row[field]
            else:
                row["_date"] = date
                records[key] = row
    stats["unique"] = len(records)
    stats["daily_files"] = len(daily)
    stats["weekly_files"] = len(weekly)
    return list(records.values()), dict(stats)


def category_path(row: dict) -> tuple[str, ...]:
    title = str(row.get("title") or "")
    low = title.casefold()
    if is_climate_policy(title):
        # The supplied taxonomy has no climate-policy branch. Its closest
        # existing route for emission governance is other environmental
        # management, never resource processing or a device technology.
        return POLICY_PATH
    raw = str(row.get("categories") or "")
    for topic in (x.strip() for x in raw.split(";")):
        if topic in TOPICS:
            if topic == "产业降碳":
                return _industrial_decarbonization(title)
            if topic == "CCUS" and re.search(r"\b(co2|carbon dioxide|carbon monoxide)\b.*\b(reduction|hydrogenation|conversion|methanation|electrolysis|to methanol|to ethanol)\b", low):
                return ("零碳产业", "物质循环", "资源加工", "有机物", "平台化工品")
            return classify_five("氢基能源" if topic == "氢氨醇" else topic, title)
    if re.search(r"\b(photovoltaic|solar cell|perovskite.silicon|solar module)\b", low):
        return classify_five("光伏", title)
    if re.search(r"\b(battery|batteries|cathode|anode|electrolyte|supercapacitor)\b", low):
        return classify_five("其他储能器件", title)
    if re.search(r"\b(hydrogen|electrolys|fuel cell|ammonia|methanol|water splitting)\b", low):
        return classify_five("氢基能源", title)
    if re.search(r"\b(carbon capture|direct air capture|co2 reduction|carbon dioxide reduction|co2 hydrogenation)\b", low):
        return classify_five("CCUS", title)
    if re.search(r"\b(plastic recycl|depolymeri|polymer recycl|waste plastic)\b", low):
        return classify_five("塑料回收", title)
    if re.search(r"\b(power grid|power system|electricity market|wind turbine|wind power)\b", low):
        return ("零碳产业", "能量转化", "能量分配与运输", "电网相关技术")
    if re.search(r"\b(ai|artificial intelligence|machine learning|deep learning|large language model|neural network)\b", low):
        return ("AI与智能科技", "AI软件层", "AI4S")
    if re.search(r"\b(data cent(?:er|re)|chip|semiconductor|transistor|robotic|robot)\b", low):
        if "data cent" in low:
            return ("AI与智能科技", "AI硬件层", "数据中心")
        if "robot" in low:
            return ("AI与智能科技", "具身智能", "硬件和控制")
        return ("AI与智能科技", "AI硬件层", "半导体")
    if re.search(r"\b(sensor|sensing|photodetector)\b", low):
        return ("通用技术", "检测和表征", "传感器")
    if re.search(r"\b(quantum comput|quantum information)\b", low):
        return ("AI与智能科技", "其它智能科技", "量子信息和量子计算")
    if re.search(r"\b(first.principles|density functional|molecular dynamics)\b", low):
        return ("通用技术", "计算仿真技术", "原子层级", "第一性原理" if "molecular dynamics" not in low else "分子动力学")
    if re.search(r"\b(catalyst|catalysis|electrocatal)\b", low):
        return ("通用技术", "材料工程", "特种功能材料", "催化材料")
    if re.search(r"\b(aviation|aircraft|railway|vehicle|shipping)\b", low):
        return ("通用技术", "通信和运输", "物质运输", "航空" if "aviation" in low or "aircraft" in low else "陆路运输")
    source = str(row.get("source") or "").casefold()
    if "biological sciences" in source or re.search(r"\b(gene|protein|cellular|patient|disease|animal|plant|neuron|brain|tumou?r|virus|ribosome|liver|ecolog|ecosystem)\b", low):
        return ("通用技术", "基础学科", "医学&生物学")
    if "phys. rev." in source or "physics letters" in source or "nuclear fusion" in source or "physical sciences : nature communications" in source:
        return ("通用技术", "基础学科", "基础物理")
    if "materials" in source or re.search(r"\b(material|polymer|nanostructure|alloy)\b", low):
        return ("通用技术", "材料工程", "其它先进材料")
    if "chem" in source or "jacs" in source or "american chemical society" in source:
        return ("通用技术", "基础学科", "基础化学")
    if "ieee transactions on smart grid" in source or "ieee transactions on power systems" in source:
        return ("零碳产业", "能量转化", "能量分配与运输", "电网相关技术")
    if "ieee transactions on energy conversion" in source:
        return ("零碳产业", "能量转化", "能源测")
    if "publication: applied energy" in source or "publication: energy policy" in source:
        return ("零碳产业", "能量转化", "能量分配与运输")
    return ("通用技术", "工艺和工程")


def make_workbook(records: list[dict], output_path: Path, template_path: Path) -> Counter:
    template = openpyxl.load_workbook(template_path)
    ws0 = template.worksheets[0]
    styles = {}
    journals = []
    publisher = ""
    for row in ws0.iter_rows(min_row=2):
        if row[0].value:
            publisher = str(row[0].value).strip()
            styles.setdefault(publisher, row)
        if row[1].value:
            journals.append((publisher, str(row[1].value).strip()))
    grouped = defaultdict(list)
    counts = Counter()
    for row in records:
        path = category_path(row)
        pub, journal = source_name(str(row.get("source") or ""))
        grouped[(path[0], pub, journal)].append((row, path))
        counts[path[0]] += 1
    for root, pub, journal in grouped:
        pair = (pub, journal)
        if pair not in journals:
            journals.append(pair)
    ordered_publishers = list(styles)
    for _, pub, _ in grouped:
        if pub not in ordered_publishers:
            ordered_publishers.append(pub)
    output = openpyxl.Workbook()
    output.remove(output.active)
    output.loaded_theme = template.loaded_theme
    for root in ROOTS:
        ws = output.create_sheet(root)
        ws.sheet_view.showGridLines = ws0.sheet_view.showGridLines
        ws.freeze_panes = ws0.freeze_panes
        for col, header in enumerate(HEADERS, 1):
            target = ws.cell(1, col, header)
            _copy_style(ws0.cell(1, min(col, 7)), target)
        for col in range(1, 7):
            letter = openpyxl.utils.get_column_letter(col)
            ws.column_dimensions[letter].width = ws0.column_dimensions[letter].width or 18
        ws.column_dimensions["G"].width = 18
        ws.column_dimensions["H"].width = ws0.column_dimensions["G"].width or 8
        ws.row_dimensions[1].height = ws0.row_dimensions[1].height or 24
        for pub in ordered_publishers:
            start_pub = ws.max_row + 1
            for p, journal in journals:
                if p != pub:
                    continue
                entries = grouped.get((root, pub, journal), [])
                start = ws.max_row + 1
                for record, path in entries or [(None, None)]:
                    values = (pub, journal,
                              record.get("title") if record else None,
                              record.get("last_author") if record else None,
                              record.get("_date") if record else None,
                              record.get("doi") or record.get("link") if record else None,
                              path[-1] if path else None, len(entries))
                    row_idx = ws.max_row + 1
                    source = styles.get(pub, styles.get("Others"))
                    for col, value in enumerate(values, 1):
                        cell = ws.cell(row_idx, col, value)
                        src_col = col if col <= 6 else 3 if col == 7 else 7
                        _copy_style(source[src_col - 1], cell)
                        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
                    ws.row_dimensions[row_idx].height = min(96, max(18, 16 * (1 + len(str(values[2] or "")) // 64)))
                if len(entries) > 1:
                    for col in (2, 8):
                        ws.merge_cells(start_row=start, start_column=col, end_row=ws.max_row, end_column=col)
            if ws.max_row > start_pub:
                ws.merge_cells(start_row=start_pub, start_column=1, end_row=ws.max_row, end_column=1)
        format_literature_worksheet(ws)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output.save(output_path)
    template.close()
    return counts
